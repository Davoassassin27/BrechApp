import os
import re
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
except ImportError:
    print("✗ Error: python-docx no está instalado")
    print("\nPara instalar:")
    print("  pip install python-docx")
    exit(1)

def parse_markdown(md_file):
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    parsed = []
    
    in_code_block = False
    code_block = []
    code_lang = ''
    
    for line in lines:
        if line.startswith('```'):
            if in_code_block:
                parsed.append(('code', code_lang, '\n'.join(code_block)))
                code_block = []
                code_lang = ''
                in_code_block = False
            else:
                in_code_block = True
                code_lang = line[3:].strip()
            continue
        
        if in_code_block:
            code_block.append(line)
            continue
        
        if line.startswith('# '):
            parsed.append(('h1', line[2:]))
        elif line.startswith('## '):
            parsed.append(('h2', line[3:]))
        elif line.startswith('### '):
            parsed.append(('h3', line[4:]))
        elif line.startswith('#### '):
            parsed.append(('h4', line[5:]))
        elif line.startswith('---'):
            parsed.append(('hr', ''))
        elif line.startswith('**') and line.endswith('**'):
            parsed.append(('bold', line[2:-2]))
        elif line.startswith('- ') or line.startswith('* '):
            parsed.append(('bullet', line[2:]))
        elif line.strip().startswith('|'):
            parsed.append(('table_row', line))
        elif line.strip() == '':
            parsed.append(('empty', ''))
        else:
            parsed.append(('text', line))
    
    return parsed

def create_docx(md_file, output_file):
    print(f"\n📄 Convirtiendo:")
    print(f"  Entrada:  {md_file}")
    print(f"  Salida:   {output_file}")
    
    doc = Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    parsed = parse_markdown(md_file)
    
    table_rows = []
    in_table = False

    for item in parsed:
        if len(item) == 2:
            item_type, content = item
            extra = None
        elif len(item) == 3:
            item_type, extra, content = item
        else:
            continue

        if item_type == 'h1':
            p = doc.add_heading(content, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        elif item_type == 'h2':
            doc.add_heading(content, level=2)

        elif item_type == 'h3':
            doc.add_heading(content, level=3)

        elif item_type == 'h4':
            doc.add_heading(content, level=4)

        elif item_type == 'hr':
            p = doc.add_paragraph()
            p.add_run('_' * 80)

        elif item_type == 'bold':
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.bold = True

        elif item_type == 'bullet':
            doc.add_paragraph(content, style='List Bullet')

        elif item_type == 'code':
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.5)

        elif item_type == 'table_row':
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(content)
        
        elif item_type == 'empty':
            if in_table and table_rows:
                create_table_from_rows(doc, table_rows)
                table_rows = []
                in_table = False
            doc.add_paragraph()
        
        elif item_type == 'text':
            if in_table and table_rows:
                create_table_from_rows(doc, table_rows)
                table_rows = []
                in_table = False
            
            content = clean_markdown_formatting(content)
            doc.add_paragraph(content)
    
    if in_table and table_rows:
        create_table_from_rows(doc, table_rows)
    
    doc.save(output_file)
    
    if os.path.exists(output_file):
        size = os.path.getsize(output_file) / 1024
        print(f"\n✓ Conversión exitosa!")
        print(f"  Archivo generado: {output_file}")
        print(f"  Tamaño: {size:.2f} KB")
        return True
    return False

def create_table_from_rows(doc, rows):
    rows = [r for r in rows if r.strip() and not r.strip().startswith('---')]
    
    if len(rows) < 2:
        return
    
    cells = [cell.strip() for cell in rows[0].split('|') if cell.strip()]
    
    if not cells:
        return
    
    table = doc.add_table(rows=1, cols=len(cells))
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    for i, cell_text in enumerate(cells):
        hdr_cells[i].text = cell_text
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    for row_text in rows[1:]:
        cells = [cell.strip() for cell in row_text.split('|') if cell.strip()]
        if cells and len(cells) == len(table.columns):
            row_cells = table.add_row().cells
            for i, cell_text in enumerate(cells):
                row_cells[i].text = cell_text

def clean_markdown_formatting(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'\[(.*?)\]\((.*?)\)', r'\1 (\2)', text)
    return text

def main():
    print("=" * 60)
    print("  Conversor Markdown → DOCX para BrechApp")
    print("  (Usando python-docx)")
    print("=" * 60)
    
    input_file = 'docs/Informe_BrechApp_MyS.md'
    output_file = 'docs/Informe_BrechApp_MyS.docx'
    
    if not os.path.exists('docs'):
        print("\n✗ Error: La carpeta 'docs' no existe")
        exit(1)
    
    if not os.path.exists(input_file):
        print(f"\n✗ Error: El archivo {input_file} no existe")
        exit(1)
    
    success = create_docx(input_file, output_file)
    
    if success:
        print("\n" + "=" * 60)
        print("  ✓ Proceso completado exitosamente")
        print("=" * 60)
        print(f"\nPuedes abrir el documento con:")
        print(f"  start {output_file}  (Windows)")
        print(f"  open {output_file}   (Mac)")
        print(f"  xdg-open {output_file}  (Linux)")
    else:
        print("\n" + "=" * 60)
        print("  ✗ La conversión falló")
        print("=" * 60)
        exit(1)

if __name__ == "__main__":
    main()
